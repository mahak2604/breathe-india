"""
Builds the formatted submission copy of the written report from docs/report/report.md.

    python docs/report/build_report.py                         -> docs/report/Breathe_India_report.docx
    powershell -File docs/report/export_report.ps1             -> fills in contents and page numbers, saves the PDF

The layout follows the MPSTME report format: a framed cover page with the NMIMS logo, a declaration,
the individual contribution table, the contents and the lists of figures and tables (pages i-iii), then
the main text with a running header and arabic page numbers from the Abstract. Word fills in the
contents, the lists and the page numbers when export_report.ps1 updates the fields.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SRC = HERE / "report.md"
OUT = HERE / "Breathe_India_report.docx"
LOGO = HERE / "assets" / "nmims_logo.png"

TITLE = "Breathe India"
SUBTITLE = "What Really Drives Air Pollution in Indian Cities?"
RUNNING_HEAD = "Breathe India – Data Wrangling Group Project Report"
COURSE = "Data Wrangling"
PROGRAMME = "B.Tech, Semester V   |   Academic Year 2026-27"
FACULTY = "Prof. Pragati Khare"
DEPARTMENT = "Department of Computer Science Engineering and Data Science, MPSTME"
DATE = "21 / 10 / 2026"
MEMBERS = [("Mahak Wadhawan", "L074"), ("Yash Shah", "L056")]
CONTRIBUTIONS = [
    ("Mahak Wadhawan", "L074",
     "Data transformation and feature engineering (star schema, joins, lag and rolling features, reshapes); "
     "the three machine-learning models and their evaluation (next-day AQI classifier, weather-normalised "
     "lockdown counterfactual, city clustering); the five-page Power BI dashboard and its DAX measures"),
    ("Yash Shah", "L056",
     "Data extraction (Kaggle download script, Open-Meteo weather and geocoding APIs, rate-limit handling and "
     "caching); data cleaning and the 23-entry cleaning log; loading to SQLite and year-partitioned Parquet "
     "with read-back validation"),
]

SERIF = "Times New Roman"
NAVY = RGBColor(0x0D, 0x36, 0x6B)
GREY = RGBColor(0x59, 0x59, 0x59)
LINK = RGBColor(0x1F, 0x5F, 0xB4)
NAVY_HEX, GRID_HEX, HEAD_FILL = "0D366B", "8C8C8C", "DCE6F2"
TEXT_W = 17.0                                     # cm between the 2.0 cm side margins
BODY_PT = 10.5
FRONT_TABLES = 2                                  # Table 1 and 2 are in the front matter

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
FIGURE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)\{([^}]*)\}")
TABLE_CAPTION = re.compile(r"^Table: (.+?)\s*\{([^}]*)\}\s*$")
XREF = re.compile(r"\{(Table|Fig):([\w-]+)\}")


# ----------------------------------------------------------------------------- low-level XML helpers
def insert_before(parent, element, successors):
    """Insert element into parent before the first child whose tag is in successors (schema order)."""
    for child in parent:
        if child.tag in [qn(s) for s in successors]:
            child.addprevious(element)
            return element
    parent.append(element)
    return element


SECTPR_ORDER = ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar",
                "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign", "noEndnote",
                "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings", "sectPrChange"]


def set_sectpr_child(section, element):
    sect_pr = section._sectPr
    tag = element.tag.split("}")[1]
    for old in sect_pr.findall(qn(f"w:{tag}")):
        sect_pr.remove(old)
    later = SECTPR_ORDER[SECTPR_ORDER.index(tag) + 1:]
    insert_before(sect_pr, element, [f"w:{t}" for t in later])


def page_border(section, style, size):
    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        for k, v in (("val", style), ("sz", str(size)), ("space", "24"), ("color", NAVY_HEX)):
            e.set(qn(f"w:{k}"), v)
        borders.append(e)
    set_sectpr_child(section, borders)


def page_numbering(section, fmt):
    num = OxmlElement("w:pgNumType")
    num.set(qn("w:fmt"), fmt)
    num.set(qn("w:start"), "1")
    set_sectpr_child(section, num)


PPR_AFTER_PBDR = ["w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                  "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
                  "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                  "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                  "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange"]


def paragraph_rule(paragraph, edge="bottom", size=6, space=3):
    pbdr = OxmlElement("w:pBdr")
    e = OxmlElement(f"w:{edge}")
    for k, v in (("val", "single"), ("sz", str(size)), ("space", str(space)), ("color", NAVY_HEX)):
        e.set(qn(f"w:{k}"), v)
    pbdr.append(e)
    insert_before(paragraph._p.get_or_add_pPr(), pbdr, PPR_AFTER_PBDR)


def add_field(paragraph, instr, shown="", **fmt):
    """A Word field (PAGE, SEQ, TOC ...). `shown` is the text displayed until Word updates it."""
    def run_with(el=None, text=None):
        run = paragraph.add_run(text) if text is not None else paragraph.add_run()
        if el is not None:
            run._r.append(el)
        style_run(run, **fmt)
        return run

    def fld(kind):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        return el

    run_with(fld("begin"))
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = f" {instr} "
    run_with(instr_el)
    run_with(fld("separate"))
    run_with(text=shown)
    run_with(fld("end"))


def style_run(run, size=None, bold=None, italic=None, color=None, font=None):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    if font is not None:
        run.font.name = font
        run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), font)
    return run


def set_cell_fill(cell, hex_fill):
    shd = OxmlElement("w:shd")
    for k, v in (("val", "clear"), ("color", "auto"), ("fill", hex_fill)):
        shd.set(qn(f"w:{k}"), v)
    insert_before(cell._tc.get_or_add_tcPr(), shd,
                  ["w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"])


def add_hyperlink(paragraph, text, url, size):
    r_id = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    for tag, val in (("w:color", "1F5FB4"), ("w:sz", str(int(size * 2)))):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        r_pr.append(el)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


# ----------------------------------------------------------------------------- text
def add_inline(paragraph, text, size=BODY_PT, bold=False, italic=False, color=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            add_inline(paragraph, part[2:-2], size, True, italic, color)
        elif part.startswith("`"):
            style_run(paragraph.add_run(part[1:-1]), size=size - 1.5, color=GREY, font="Consolas")
        elif part.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            add_hyperlink(paragraph, label, url, size)
        elif part.startswith("*"):
            add_inline(paragraph, part[1:-1], size, bold, True, color)
        else:
            style_run(paragraph.add_run(part), size=size, bold=bold or None, italic=italic or None, color=color)


def body_para(container, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=BODY_PT, after=5, keep=False):
    p = container.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    add_inline(p, text, size)
    return p


def centred(doc, text, size, bold=False, italic=False, color=None, before=0, after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
    p.paragraph_format.line_spacing = 1.0
    style_run(p.add_run(text), size=size, bold=bold, italic=italic, color=color)
    return p


def numbered(doc, number, text, size=BODY_PT, hang=0.9, after=4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.left_indent, pf.first_line_indent, pf.space_after = Cm(hang), Cm(-hang + 0.25), Pt(after)
    pf.tab_stops.add_tab_stop(Cm(hang))
    style_run(p.add_run(f"{number}\t"), size=size)
    add_inline(p, text, size)
    return p


def caption(container, kind, number, text, size=9.5, before=3, after=9, keep=False):
    """'Fig. 3. text' or 'Table 3. text'; the SEQ field lets Word build the lists of figures and tables."""
    p = container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.keep_with_next, pf.line_spacing = Pt(before), Pt(after), keep, 1.0
    style_run(p.add_run("Fig. " if kind == "Figure" else "Table "), size=size, bold=True)
    add_field(p, f"SEQ {kind} \\* ARABIC", str(number), size=size, bold=True)
    style_run(p.add_run(". "), size=size, bold=True)
    add_inline(p, text, size)
    return p


# ----------------------------------------------------------------------------- tables and figures
def format_table(table, widths, borders=True):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 567)))
    tbl_w.set(qn("w:type"), "dxa")
    if borders:
        el = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", GRID_HEX)):
                e.set(qn(f"w:{k}"), v)
            el.append(e)
        insert_before(tbl_pr, el, ["w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"])
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    insert_before(tbl_pr, layout, ["w:tblCellMar", "w:tblLook"])
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", 35), ("left", 85), ("bottom", 35), ("right", 85)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    insert_before(tbl_pr, mar, ["w:tblLook"])
    for col, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(int(w * 567)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = Cm(w)


def row_flags(row, header=False):
    tr_pr = row._tr.get_or_add_trPr()
    flags = ["cantSplit"] + (["tblHeader"] if header else [])
    for flag in flags:
        el = OxmlElement(f"w:{flag}")
        successors = {"cantSplit": ["w:trHeight", "w:tblHeader", "w:tblCellSpacing", "w:jc", "w:hidden"],
                      "tblHeader": ["w:tblCellSpacing", "w:jc", "w:hidden"]}[flag]
        insert_before(tr_pr, el, successors)


def data_table(doc, rows, widths, align=None, size=9.5):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_table(table, widths)
    aligns = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT}
    for i, row in enumerate(rows):
        row_flags(table.rows[i], header=(i == 0))
        for j, text in enumerate(row):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER if i == 0 else WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else aligns[(align or "l" * len(row))[j]]
            add_inline(p, text, size=size, bold=(i == 0))
            if i == 0:
                set_cell_fill(cell, HEAD_FILL)
            # a short table stays on one page with its caption; a long one may break between rows
            p.paragraph_format.keep_with_next = len(rows) <= 12 and i < len(rows) - 1
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    spacer.paragraph_format.line_spacing = 0.8
    return table


def picture(container, path, width, after=0):
    p = container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width))
    return p


def figures_row(doc, figs, numbers):
    """One figure centred, or two side by side, each with its caption underneath."""
    if len(figs) == 1:
        (cap, path, attrs), = figs
        picture(doc, path, float(attrs["w"]))
        caption(doc, "Figure", numbers[0], cap)
        return
    widths = [float(a["w"]) + 0.3 for _, _, a in figs]
    table = doc.add_table(rows=1, cols=len(figs))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_table(table, widths, borders=False)
    row_flags(table.rows[0])
    for cell, (cap, path, attrs), n in zip(table.rows[0].cells, figs, numbers):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        first = cell.paragraphs[0]
        first.alignment = WD_ALIGN_PARAGRAPH.CENTER
        first.paragraph_format.space_before = Pt(6)
        first.add_run().add_picture(str(path), width=Cm(float(attrs["w"])))
        caption(cell, "Figure", n, cap, size=9, after=2)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)
    spacer.paragraph_format.line_spacing = 0.8


def parse_attrs(text):
    attrs = {}
    for token in text.split():
        if token.startswith("#"):
            attrs["label"] = token[1:]
        elif "=" in token:
            k, v = token.split("=", 1)
            attrs[k] = v
    return attrs


def table_rows(lines):
    rows = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-+:?", c) for c in cells):
            rows.append(cells)
    return rows


# ----------------------------------------------------------------------------- styles, sections, header
def set_style_font(style, size, bold=False, color=None, italic=False):
    style.font.size, style.font.bold, style.font.italic = Pt(size), bold, italic
    if color is not None:
        style.font.color.rgb = color
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), SERIF)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    set_style_font(normal, BODY_PT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.0
    for name, size, before, after in (("Heading 1", 14, 12, 5), ("Heading 2", 12, 8, 3)):
        st = doc.styles[name]
        set_style_font(st, size, bold=True, color=NAVY)
        pf = st.paragraph_format
        pf.space_before, pf.space_after, pf.keep_with_next, pf.line_spacing = Pt(before), Pt(after), True, 1.0
    lb = doc.styles["List Bullet"]
    lb.paragraph_format.space_after = Pt(3)


def page_setup(section):
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    section.top_margin, section.bottom_margin = Cm(2.4), Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.0)
    section.header_distance = section.footer_distance = Cm(1.2)


def header_footer(section):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
    hp.paragraph_format.space_after = Pt(0)
    hp.add_run().add_picture(str(LOGO), height=Cm(0.8))
    style_run(hp.add_run("\t" + RUNNING_HEAD), size=9, italic=True, color=GREY)
    paragraph_rule(hp, "bottom")
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(fp.add_run(f"{TITLE}   |   Page "), size=9, color=GREY)
    add_field(fp, "PAGE", "1", size=9, color=GREY)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ----------------------------------------------------------------------------- front matter
def cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.add_run().add_picture(str(LOGO), width=Cm(4.0))
    centred(doc, "SVKM's NMIMS (Deemed-to-be University)", 15, bold=True, before=8)
    centred(doc, "Mukesh Patel School of Technology Management & Engineering", 12.5, bold=True, before=3)
    centred(doc, "Vile Parle (West), Mumbai - 400 056", 10.5, before=2)
    centred(doc, "A Data Wrangling Group Project Report on", 12, italic=True, before=44)
    centred(doc, TITLE, 32, bold=True, color=NAVY, before=6)
    centred(doc, SUBTITLE, 15, bold=True, before=4)
    centred(doc, f"Course: {COURSE}", 12, before=34)
    centred(doc, PROGRAMME, 12, before=3)
    centred(doc, "Submitted by", 12.5, bold=True, before=40, after=4)

    table = doc.add_table(rows=len(MEMBERS) + 1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_table(table, [5.2, 3.0], borders=False)
    for i, (name, roll) in enumerate([("Name", "Roll No.")] + MEMBERS):
        for cell, text in zip(table.rows[i].cells, (name, roll)):
            cp = cell.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(1)
            style_run(cp.add_run(text), size=12, bold=(i == 0))
        if i == 0:
            paragraph_rule(table.rows[0].cells[0].paragraphs[0], size=4, space=1)
            paragraph_rule(table.rows[0].cells[1].paragraphs[0], size=4, space=1)

    centred(doc, "Under the guidance of", 12, italic=True, before=40)
    centred(doc, FACULTY, 14, bold=True, before=3)
    centred(doc, DEPARTMENT, 10.5, before=3)
    centred(doc, f"Date of Submission: {DATE}", 11.5, before=46)


def front_matter(doc):
    doc.add_heading("Declaration", level=1)
    body_para(doc, "We, the undersigned students of B.Tech Semester V, Mukesh Patel School of Technology Management & "
                   "Engineering, SVKM's NMIMS, Mumbai, declare that this report titled "
                   f"**\"{TITLE}: {SUBTITLE}\"** is a record of the work carried out by us as part of the "
                   f"{COURSE} group project during the academic year 2026-27. The data used is publicly available, "
                   "and all sources of data, information and software have been acknowledged in the References "
                   "section.")
    body_para(doc, "Place: Mumbai", align=WD_ALIGN_PARAGRAPH.LEFT, after=2)
    body_para(doc, f"Date: {DATE}", align=WD_ALIGN_PARAGRAPH.LEFT, after=10)
    caption(doc, "Table", 1, "Declaration signatures", keep=True, after=4)
    rows = [("Name", "Roll No.", "Signature")] + [(n, r, "") for n, r in MEMBERS]
    table = data_table(doc, rows, [6.0, 3.0, 6.0], align="lcl", size=10.5)
    for row in table.rows[1:]:
        row.height_rule, row.height = WD_ROW_HEIGHT_RULE.AT_LEAST, Cm(1.0)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    doc.add_heading("Individual Contribution", level=1)
    body_para(doc, "The table below lists the primary responsibility of each member. Both members worked together "
                   "on the insights, this report and the viva presentation, and reviewed each other's parts of "
                   "the pipeline.")
    caption(doc, "Table", 2, "Individual contribution of team members", keep=True, after=4)
    data_table(doc, [("Name", "Roll No.", "Specific tasks handled")] + CONTRIBUTIONS + [
        ("Both members", "L074, L056", "Insights and their verification; this written report; the viva presentation")],
        [3.4, 2.0, 11.6], align="lcl", size=10)
    page_break(doc)

    doc.add_heading("Table of Contents", level=1)
    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-2" \\h \\z \\u', "Update the fields to build the table of contents.")
    page_break(doc)

    doc.add_heading("List of Figures", level=1)
    add_field(doc.add_paragraph(), 'TOC \\h \\z \\c "Figure"', "Update the fields to build the list of figures.")
    doc.add_heading("List of Tables", level=1)
    add_field(doc.add_paragraph(), 'TOC \\h \\z \\c "Table"', "Update the fields to build the list of tables.")


# ----------------------------------------------------------------------------- main text
def number_items(lines):
    """First pass: give every labelled table and figure its number, in reading order."""
    numbers, t, f = {}, FRONT_TABLES, 0
    for line in lines:
        m = TABLE_CAPTION.match(line)
        if m:
            t += 1
            numbers[("Table", parse_attrs(m.group(2)).get("label"))] = t
        for fig in FIGURE.finditer(line):
            f += 1
            numbers[("Fig", parse_attrs(fig.group(3)).get("label"))] = f
    return numbers


def main_text(doc, lines, numbers):
    def xref(text):
        return XREF.sub(lambda m: f"{'Table ' if m.group(1) == 'Table' else 'Fig. '}{numbers[(m.group(1), m.group(2))]}",
                        text)

    pending_table, section = None, ""
    i = next(k for k, l in enumerate(lines) if l.startswith("## "))
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("<!--"):
            if not stripped.endswith("-->"):                       # a multi-line comment: skip it
                while not lines[i].strip().endswith("-->"):
                    i += 1
            elif "pagebreak" in stripped:
                page_break(doc)
            elif "include:" in stripped:
                path = (HERE / stripped.split("include:")[1].replace("-->", "").strip()).resolve()
                inc = [l for l in path.read_text(encoding="utf-8").splitlines() if l.startswith("|")]
                rows = [[c.replace("ug/m3", "µg/m³") for c in r] for r in table_rows(inc)]
                add_table(doc, rows, pending_table, numbers)
                pending_table = None
            i += 1
            continue
        if line.startswith("## "):
            section = line[3:]
            doc.add_heading(section, level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif TABLE_CAPTION.match(line):
            m = TABLE_CAPTION.match(line)
            pending_table = (m.group(1), parse_attrs(m.group(2)))
        elif line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(xref(lines[i]))
                i += 1
            add_table(doc, table_rows(block), pending_table, numbers)
            pending_table = None
            continue
        elif FIGURE.match(stripped):
            figs = [(xref(m.group(1)), (HERE / m.group(2)).resolve(), parse_attrs(m.group(3)))
                    for m in FIGURE.finditer(stripped)]
            figures_row(doc, figs, [numbers[("Fig", a.get("label"))] for _, _, a in figs])
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, xref(line[2:]))
        elif re.match(r"\d+\. ", line):
            n, text = line.split(". ", 1)
            if section == "References":
                numbered(doc, f"[{n}]", text, size=10, hang=0.9, after=3)
            else:
                numbered(doc, f"{n}.", xref(text))
        elif line.startswith("**Keywords:**"):
            body_para(doc, line, align=WD_ALIGN_PARAGRAPH.LEFT)
        else:
            keep = line.endswith(":")
            body_para(doc, xref(line), keep=keep)
        i += 1


def add_table(doc, rows, pending, numbers):
    text, attrs = pending
    widths = [float(w) for w in attrs["widths"].split(",")]
    widths = [w * TEXT_W / sum(widths) for w in widths]
    assert len(widths) == len(rows[0]), (text, len(widths), len(rows[0]))
    caption(doc, "Table", numbers[("Table", attrs.get("label"))], text, keep=True, before=6, after=4)
    data_table(doc, rows, widths, attrs.get("align"), size=float(attrs.get("size", 9)))


# ----------------------------------------------------------------------------- build
def build():
    doc = Document()
    setup_styles(doc)
    cover(doc)
    doc.add_section(WD_SECTION.NEW_PAGE)
    front_matter(doc)
    doc.add_section(WD_SECTION.NEW_PAGE)
    lines = SRC.read_text(encoding="utf-8").splitlines()
    main_text(doc, lines, number_items(lines))

    cover_s, front_s, body_s = doc.sections
    for s in doc.sections:
        page_setup(s)
    page_border(cover_s, "double", 6)
    page_border(front_s, "single", 6)
    page_border(body_s, "single", 6)
    page_numbering(front_s, "lowerRoman")
    page_numbering(body_s, "decimal")
    header_footer(front_s)                      # the main text inherits it; the cover has none

    props = doc.core_properties
    props.title = f"{TITLE}: {SUBTITLE}"
    props.subject = f"{COURSE} group project report"
    props.author = ", ".join(n for n, _ in MEMBERS)
    props.keywords = "data wrangling, air quality, CPCB, Power BI"
    doc.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    build()
